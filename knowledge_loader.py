"""อ่านและตรวจข้อมูลร้านจากไฟล์ภายนอก โดยไม่เรียก LLM."""

import math
import re
from dataclasses import dataclass
from numbers import Real
from pathlib import Path

from config import MENU_FILE, MENU_SHEET, STORE_INFO_FILE, STORE_POLICIES_FILE


class KnowledgeDataError(ValueError):
    """ข้อมูลไม่ครบหรือมีค่าที่ระบบไม่รองรับ."""


@dataclass(frozen=True)
class Document:
    # เก็บแหล่งที่มาไว้สำหรับอ้างอิงผลค้นหา
    id: str
    title: str
    text: str
    source: str


@dataclass(frozen=True)
class KnowledgeBase:
    menu: tuple[dict, ...]
    documents: tuple[Document, ...]


REQUIRED_COLUMNS = (
    "id", "name", "category", "price", "temperatures", "caffeine_status",
    "sweetness", "sweetness_options", "description",
)
TEMPERATURE_VALUES = {
    "ร้อน": ("ร้อน",),
    "เย็น": ("เย็น",),
    "ร้อนและเย็น": ("ร้อน", "เย็น"),
    "ไม่เกี่ยวข้อง": (),
    "ยังไม่ทราบ": None,
}
CAFFEINE_VALUES = {"มี", "ไม่มี", "ยังไม่ทราบ"}


def load_menu(path: Path = MENU_FILE, sheet_name: str = MENU_SHEET) -> tuple[dict, ...]:
    """อ่านเมนูและหยุดพร้อมบอกแถวที่ผิด หากข้อมูลใช้ไม่ได้."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise KnowledgeDataError("ติดตั้งก่อน: python -m pip install openpyxl") from exc

    path = Path(path)
    if not path.is_file():
        raise KnowledgeDataError(f"ไม่พบไฟล์เมนู: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise KnowledgeDataError(f"เปิด Excel ไม่ได้: {path.name} ({exc})") from exc

    try:
        if sheet_name not in workbook.sheetnames:
            raise KnowledgeDataError(f"{path.name}: ไม่พบ sheet '{sheet_name}'")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows()
        header_cells = next(rows, ())
        if not header_cells:
            raise KnowledgeDataError(f"{path.name}: ไม่มีหัวคอลัมน์")
        headers = [str(c.value).strip() if c.value is not None else "" for c in header_cells]
        named_headers = [h for h in headers if h]
        if len(named_headers) != len(set(named_headers)):
            raise KnowledgeDataError(f"{path.name}: หัวคอลัมน์ซ้ำกัน")
        missing = set(REQUIRED_COLUMNS) - set(headers)
        if missing:
            raise KnowledgeDataError(f"{path.name}: ขาดคอลัมน์ {', '.join(sorted(missing))}")
        indices = {h: headers.index(h) for h in REQUIRED_COLUMNS}
        items, seen_ids = [], set()

        for row_number, cells in enumerate(rows, start=2):
            selected = {h: cells[i] for h, i in indices.items()}
            if all(c.value is None or str(c.value).strip() == "" for c in selected.values()):
                continue  # ข้ามแถวว่าง ไม่ข้ามแถวที่กรอกเพียงบางช่อง
            where = f"{path.name} แถว {row_number}"
            # รับข้อมูลจริง ไม่ใช้สูตรหรือค่าคำนวณที่อาจยังไม่อัปเดต
            if any(c.data_type == "f" for c in selected.values()):
                raise KnowledgeDataError(f"{where}: ใช้ค่าข้อมูลแทนสูตร Excel")
            item = {h: c.value for h, c in selected.items()}
            for h in REQUIRED_COLUMNS:
                if h == "price":
                    continue
                if not isinstance(item[h], str) or not item[h].strip():
                    raise KnowledgeDataError(f"{where}: {h} ต้องเป็นข้อความและไม่ว่าง")
                item[h] = item[h].strip()
            if item["id"] in seen_ids:
                raise KnowledgeDataError(f"{where}: รหัสเมนูซ้ำ {item['id']}")
            seen_ids.add(item["id"])

            price = item["price"]
            if isinstance(price, bool) or not isinstance(price, Real) or not math.isfinite(price) or price < 0:
                raise KnowledgeDataError(f"{where}: price ต้องเป็นตัวเลขตั้งแต่ 0 ขึ้นไป")
            if item["temperatures"] not in TEMPERATURE_VALUES:
                raise KnowledgeDataError(f"{where}: temperatures ใช้ค่าที่กำหนดในข้อมูลเมนู")
            if item["caffeine_status"] not in CAFFEINE_VALUES:
                raise KnowledgeDataError(f"{where}: caffeine_status ต้องเป็น มี / ไม่มี / ยังไม่ทราบ")

            # None คือไม่ทราบ ไม่ใช่ไม่มี; () คือไม่เกี่ยวข้องกับอุณหภูมิเครื่องดื่ม
            item["temperature_options"] = TEMPERATURE_VALUES[item["temperatures"]]
            item["source"] = f"{path.name} / {sheet_name} / {item['id']}"
            items.append(item)

        if not items:
            raise KnowledgeDataError(f"{path.name}: ไม่มีรายการเมนู")
        return tuple(items)
    finally:
        workbook.close()


def load_document(path: Path) -> tuple[Document, ...]:
    """แบ่ง Markdown ตามหัวข้อระดับ ## และรักษาตารางในแต่ละหัวข้อ."""
    path = Path(path)
    try:
        text = path.read_text(encoding="utf-8-sig").strip()
    except (OSError, UnicodeError) as exc:
        raise KnowledgeDataError(f"อ่านเอกสารไม่ได้: {path} ({exc})") from exc
    if not text:
        raise KnowledgeDataError(f"{path.name}: เอกสารว่าง")

    headings = list(re.finditer(r"^##[ \t]+(.+)$", text, flags=re.MULTILINE))
    if not headings:
        raise KnowledgeDataError(f"{path.name}: ต้องมีหัวข้อระดับ ##")
    documents, seen_ids = [], set()
    for n, heading in enumerate(headings):
        title = heading.group(1).strip()
        end = headings[n + 1].start() if n + 1 < len(headings) else len(text)
        body = text[heading.end():end].strip()
        if not body:
            raise KnowledgeDataError(f"{path.name}: หัวข้อ '{title}' ไม่มีเนื้อหา")
        # ใช้รหัส STORE/POLICY ที่มีอยู่ หรือสร้างจากชื่อไฟล์และลำดับ
        explicit_id = re.match(r"((?:STORE|POLICY)-\d+)\b", title)
        document_id = explicit_id.group(1) if explicit_id else f"{path.stem}-{n + 1:03d}"
        if document_id in seen_ids:
            raise KnowledgeDataError(f"{path.name}: รหัสหัวข้อซ้ำ {document_id}")
        seen_ids.add(document_id)
        documents.append(Document(document_id, title, f"{title}\n{body}", path.name))
    return tuple(documents)


def load_knowledge_base() -> KnowledgeBase:
    """รวมเมนู ข้อมูลร้าน และนโยบาย ให้ส่วนค้นหานำไปใช้ต่อ."""
    menu = load_menu()
    documents = load_document(STORE_INFO_FILE) + load_document(STORE_POLICIES_FILE)
    ids = [d.id for d in documents]
    if len(ids) != len(set(ids)):
        raise KnowledgeDataError("รหัสหัวข้อซ้ำข้ามไฟล์เอกสาร")
    return KnowledgeBase(menu, documents)


if __name__ == "__main__":
    try:
        knowledge = load_knowledge_base()
    except KnowledgeDataError as exc:
        print(f"โหลดข้อมูลไม่สำเร็จ: {exc}")
        raise SystemExit(1)
    print(f"โหลดสำเร็จ: {len(knowledge.menu)} เมนู / {len(knowledge.documents)} หัวข้อเอกสาร")
    for item in knowledge.menu[:3]:
        print(f"  {item['id']} {item['name']} — {item['price']:g} บาท")
    print("ยังไม่ทำ embedding และยังไม่เรียก API")
